import io
import json
import os
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.ingestion import Store, clean_text, extract_text
from app.main import app


def docx_bytes(text="熟悉 Python 和 MySQL，负责后端服务开发和数据库设计，参与多个校园项目。", table=True):
    document = Document()
    document.add_paragraph(text)
    if table:
        cells = document.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text, cells[1].text = "学校", "计算机科学本科"
        document.add_paragraph("项目成果：完成接口交付。")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_bytes(text=None, encrypted=False, pages=1):
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=600, height=800)
        if text:
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                     NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({
                NameObject("/F1"): writer._add_object(font)})})
            content = DecodedStreamObject()
            content.set_data(f"BT /F1 10 Tf 30 700 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(content)
    if encrypted:
        writer.encrypt("secret")
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


class IngestApiTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.environment = patch.dict(os.environ, {"DATA_DIR": self.directory.name, "MAX_UPLOAD_MB": "15",
                                                   "OCR_MAX_PAGES": "10"})
        self.environment.start()
        self.client_context = TestClient(app)
        self.client = self.client_context.__enter__()

    def tearDown(self):
        self.client_context.__exit__(None, None, None)
        self.environment.stop()
        self.directory.cleanup()

    def upload(self, content, name="resume.docx"):
        response = self.client.post("/api/resumes/batch", files=[("files[]", (name, content))])
        self.assertEqual(response.status_code, 202, response.text)
        return response.json()["task_id"]

    def wait(self, task_id):
        for _ in range(300):
            response = self.client.get(f"/api/tasks/{task_id}")
            self.assertEqual(response.status_code, 200)
            data = response.json()
            if data["status"] in {"succeeded", "failed"}:
                return data
            time.sleep(0.01)
        self.fail("task never completed")

    def assert_error(self, response, status, code):
        self.assertEqual(response.status_code, status, response.text)
        error = response.json()["error"]
        self.assertEqual(error["code"], code)
        self.assertTrue(error["request_id"])
        self.assertTrue(error["message"])

    def test_docx_roundtrip_order_and_dedup(self):
        content = docx_bytes()
        with patch("app.ingestion.extract_text", wraps=extract_text) as extract:
            first = self.upload(content, "../../张三.docx")
            second = self.upload(content, "renamed.pdf")
            result, duplicate = self.wait(first), self.wait(second)
            self.assertEqual(extract.call_count, 1)
        self.assertEqual(result["status"], "succeeded")
        self.assertEqual(result["progress"], 100)
        item = result["items"][0]
        self.assertEqual(item["name"], "张三.docx")
        self.assertEqual(item["resume_id"], duplicate["items"][0]["resume_id"])
        text = self.client.get(f'/api/resumes/{item["resume_id"]}/text').json()["text"]
        self.assertIn("学校 | 计算机科学本科", text)
        self.assertLess(text.index("学校"), text.index("项目成果"))
        self.assertEqual(item["chars"], len(text))
        self.assertEqual(len(list((Path(self.directory.name) / "uploads").iterdir())), 1)
        self.assertFalse(list(Path(self.directory.name).glob("*.tmp")))

    def test_frontend_field_roundtrip_and_legacy_dedup(self):
        content = docx_bytes()
        response = self.client.post("/api/resumes/batch", files=[("files", ("前端.docx", content))])
        self.assertEqual(response.status_code, 202, response.text)
        result = self.wait(response.json()["task_id"])
        self.assertEqual(result["status"], "succeeded")
        self.assertIsNone(result["error"])
        resume_id = result["items"][0]["resume_id"]
        self.assertEqual(self.wait(self.upload(content))["items"][0]["resume_id"], resume_id)
        text = self.client.get(f"/api/resumes/{resume_id}/text")
        self.assertEqual(text.status_code, 200)
        self.assertIn("Python", text.json()["text"])

    def test_both_upload_fields_share_validation_limits(self):
        response = self.client.post("/api/resumes/batch", files=[
            ("files", ("new.docx", docx_bytes("候选人一，负责后端服务。"))),
            ("files[]", ("legacy.docx", docx_bytes("候选人二，负责数据库。")))])
        self.assertEqual(response.status_code, 202, response.text)
        self.assertEqual(len(self.wait(response.json()["task_id"])["items"]), 2)
        entries = [("files" if index < 11 else "files[]", (f"{index}.docx", b"x")) for index in range(21)]
        self.assert_error(self.client.post("/api/resumes/batch", files=entries), 400, "TOO_MANY_RESUMES")
        self.assert_error(self.client.post("/api/resumes/batch"), 422, "VALIDATION_ERROR")
        self.assert_error(self.client.post("/api/resumes/batch", files=[("files", ("empty.docx", b""))]),
                          400, "INVALID_INPUT")

    def test_polling_never_exposes_parsing_items(self):
        started, release = threading.Event(), threading.Event()

        def blocked_extract(*args):
            started.set()
            if not release.wait(10):
                raise RuntimeError("test extraction was not released")
            return extract_text(*args)

        with patch("app.ingestion.extract_text", side_effect=blocked_extract):
            try:
                first = self.upload(docx_bytes("第一份简历，熟悉Python和MySQL，负责后端开发与测试。"))
                self.assertTrue(started.wait(3))
                second = self.upload(docx_bytes("第二份简历，熟悉Java和MySQL，负责后端服务开发。"))
                for identifier, expected in ((first, "running"), (second, "queued")):
                    result = self.client.get(f"/api/tasks/{identifier}").json()
                    self.assertEqual(result["status"], expected)
                    self.assertEqual(result["items"], [])
                    self.assertIsNone(result["error"])
                    with app.state.store.connect() as db:
                        self.assertGreater(db.execute("SELECT count(*) FROM resumes WHERE status='parsing'").fetchone()[0], 0)
            finally:
                release.set()
            for identifier in (first, second):
                result = self.wait(identifier)
                self.assertEqual(result["status"], "succeeded")
                self.assertTrue(result["items"])
                self.assertTrue(all(item["status"] in {"ready", "need_review", "failed"} for item in result["items"]))

    def test_restart_error_is_frontend_renderable_string(self):
        with app.state.store.connect() as db:
            db.execute("INSERT INTO resumes (id,name,sha256,file_path,source,created_at) VALUES ('interrupted','x','x','x','docx','now')")
            db.execute("INSERT INTO tasks (id,kind,status,payload,created_at,updated_at) VALUES "
                       "('interrupted','ingest','running',?,'now','now')", (json.dumps({"resume_ids": ["interrupted"]}),))
        app.state.store.initialize()
        result = self.client.get("/api/tasks/interrupted").json()
        self.assertEqual(result["status"], "failed")
        self.assertIsInstance(result["error"], str)
        self.assertIn("服务重启", result["error"])
        self.assertEqual(result["error_code"], "TASK_INTERRUPTED")
        self.assertEqual(result["items"][0]["status"], "failed")

    def test_text_pdf_magic_over_extension(self):
        content = pdf_bytes("Python MySQL backend engineer with experience in database design and REST APIs. " * 3)
        result = self.wait(self.upload(content, "resume.doc"))
        self.assertEqual(result["status"], "succeeded")
        item = result["items"][0]
        text = self.client.get(f'/api/resumes/{item["resume_id"]}/text').json()["text"]
        self.assertIn("Python MySQL", text)

    def test_short_and_empty_docx(self):
        result = self.wait(self.upload(docx_bytes("短简历", table=False)))
        self.assertEqual(result["items"][0]["status"], "need_review")
        result = self.wait(self.upload(docx_bytes("", table=False)))
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["items"][0]["error"]["code"], "OCR_FAILED")

    def test_invalid_input_atomic_validation(self):
        for content, status, code in [(b"", 400, "INVALID_INPUT"), (b"old doc", 415, "UNSUPPORTED_FILE"),
                                      (b"x" * (15 * 1024 * 1024 + 1), 413, "FILE_TOO_LARGE")]:
            response = self.client.post("/api/resumes/batch", files=[("files[]", ("test.pdf", content))])
            self.assert_error(response, status, code)
        response = self.client.post("/api/resumes/batch", files=[
            ("files[]", ("valid.docx", docx_bytes())), ("files[]", ("bad.pdf", b"bad"))])
        self.assert_error(response, 415, "UNSUPPORTED_FILE")
        with app.state.store.connect() as db:
            self.assertEqual(db.execute("SELECT count(*) FROM resumes").fetchone()[0], 0)
            self.assertEqual(db.execute("SELECT count(*) FROM tasks").fetchone()[0], 0)
        self.assertFalse(list(Path(self.directory.name).glob("*.tmp")))

    def test_batch_limit_and_twenty_files(self):
        files = [("files[]", (f"{index}.docx", docx_bytes(f"候选人{index}，熟悉Python和MySQL，负责后端开发与数据库设计，完成接口上线和测试。")))
                 for index in range(21)]
        self.assert_error(self.client.post("/api/resumes/batch", files=files), 400, "TOO_MANY_RESUMES")
        response = self.client.post("/api/resumes/batch", files=files[:20])
        self.assertEqual(response.status_code, 202)
        result = self.wait(response.json()["task_id"])
        self.assertEqual(result["status"], "succeeded")
        self.assertEqual(len(result["items"]), 20)
        self.assertTrue(all(item["status"] == "ready" for item in result["items"]))

    def test_ocr_unavailable_encrypted_corrupt_and_page_limit(self):
        for content, code in [(pdf_bytes(), "OCR_UNAVAILABLE"),
                              (b"\x89PNG\r\n\x1a\nimage", "OCR_UNAVAILABLE"),
                              (b"\xff\xd8\xffimage", "OCR_UNAVAILABLE"),
                              (pdf_bytes(encrypted=True), "UNSUPPORTED_FILE"),
                              (pdf_bytes(pages=11), "UNSUPPORTED_FILE"),
                              (b"%PDF-broken", "UNSUPPORTED_FILE")]:
            result = self.wait(self.upload(content))
            self.assertEqual(result["status"], "failed")
            item = result["items"][0]
            self.assertEqual(item["status"], "failed")
            self.assertEqual(item["error"]["code"], code)
            self.assertIsInstance(result["error"], str)
            self.assertEqual(result["error"], item["error"]["message"])
            self.assertEqual(result["error_code"], code)
            response = self.client.get(f'/api/resumes/{item["resume_id"]}/text')
            self.assert_error(response, 503 if code == "OCR_UNAVAILABLE" else 415, code)

    def test_partial_success_and_retry_failed(self):
        response = self.client.post("/api/resumes/batch", files=[
            ("files[]", ("ok.docx", docx_bytes())), ("files[]", ("scan.pdf", pdf_bytes()))])
        result = self.wait(response.json()["task_id"])
        self.assertEqual(result["status"], "succeeded")
        self.assertEqual([item["status"] for item in result["items"]], ["ready", "failed"])
        old_id = result["items"][1]["resume_id"]
        retried = self.wait(self.upload(pdf_bytes()))
        self.assertEqual(retried["items"][0]["resume_id"], old_id)
        self.assertEqual(retried["status"], "failed")

    def test_unexpected_parser_failure_finishes_rows(self):
        with patch("app.ingestion.extract_text", side_effect=RuntimeError("unexpected")):
            result = self.wait(self.upload(docx_bytes()))
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["items"][0]["status"], "failed")
        self.assertIsInstance(result["error"], str)
        self.assertEqual(result["error_code"], "INGEST_FAILED")

    def test_missing_ids_docs_disabled_and_not_ready(self):
        self.assert_error(self.client.get("/api/tasks/nope"), 404, "TASK_NOT_FOUND")
        self.assert_error(self.client.get("/api/resumes/nope/text"), 404, "RESUME_NOT_FOUND")
        for path in ("/docs", "/redoc", "/openapi.json"):
            self.assertEqual(self.client.get(path).status_code, 404)
        with app.state.store.connect() as db:
            db.execute("INSERT INTO resumes (id,name,sha256,file_path,source,created_at) VALUES ('pending','x','x','x','docx','now')")
        self.assert_error(self.client.get("/api/resumes/pending/text"), 409, "RESUME_NOT_READY")


class RecoveryAndCleaningTests(unittest.TestCase):
    def test_upload_schema_uses_binary_file_picker(self):
        schema = app.openapi()
        body = schema["paths"]["/api/resumes/batch"]["post"]["requestBody"]
        reference = body["content"]["multipart/form-data"]["schema"]["$ref"]
        fields = schema["components"]["schemas"][reference.rsplit("/", 1)[-1]]
        for name in ("files", "files[]"):
            upload = fields["properties"][name]
            self.assertEqual(upload["type"], "array")
            self.assertEqual(upload["items"]["format"], "binary")

    def test_recovery_and_persistence(self):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory))
            store.initialize()
            with store.connect() as db:
                for state in ("queued", "running", "succeeded"):
                    db.execute("INSERT INTO tasks (id,kind,status,created_at,updated_at) VALUES (?,'ingest',?,'now','now')", (state, state))
                for state in ("parsing", "ready"):
                    db.execute("INSERT INTO resumes (id,name,sha256,file_path,source,status,created_at) VALUES (?,?,?,'x','docx',?,'now')",
                               (state, state, state, state))
            restarted = Store(Path(directory))
            restarted.initialize()
            self.assertEqual(restarted.get("tasks", "queued")["status"], "failed")
            self.assertEqual(restarted.get("tasks", "running")["status"], "failed")
            self.assertEqual(restarted.get("tasks", "succeeded")["status"], "succeeded")
            self.assertEqual(restarted.get("resumes", "parsing")["status"], "failed")
            self.assertEqual(restarted.get("resumes", "ready")["status"], "ready")
            self.assertIn("服务重启", json.loads(restarted.get("tasks", "queued")["error"])["message"])

    def test_conservative_cleaning(self):
        cleaned = clean_text("\ufeff标题\n第 1 页\n页眉\n\n\nPython\n2025-2026\n页眉\n学校\u3000本科\n页眉\n内部资料\n个人网站 www.example.com 是我的作品集\n")
        self.assertEqual(cleaned, "标题\n\nPython\n2025-2026\n学校 本科\n个人网站 www.example.com 是我的作品集")


if __name__ == "__main__":
    unittest.main()
