import asyncio
import io
import json
import os
import tempfile
import time
import unittest
from unittest.mock import AsyncMock, patch

from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.matching import LlmError, cache_key, education_score, jd_terms, run_diagnose
from app.mock_data import MOCK_RESPONSES
from app.schemas import DiagnoseResponse

JD = "本科及以上学历，要求熟悉Python和MySQL，掌握FastAPI接口开发，有团队协作和单元测试经验。"
TEXT = "郑州大学计算机本科，熟悉Python和MySQL，使用FastAPI开发后端接口，负责单元测试与团队协作。"


class MatchingTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.environment = patch.dict(os.environ, {"DATA_DIR": self.directory.name, "MATCH_LLM_TOP_K": "0",
                                                   "MATCH_CONCURRENCY": "3", "OCR_PROVIDER": "stub"})
        self.environment.start()
        self.context = TestClient(app)
        self.client = self.context.__enter__()
        with app.state.store.connect() as db:
            for identifier, text in (("a", TEXT), ("b", TEXT), ("c", "高中毕业，热爱摄影和旅游，曾参加校园志愿活动，负责活动组织与现场接待。")):
                db.execute("INSERT INTO resumes (id,name,sha256,file_path,source,text,chars,status,created_at) "
                           "VALUES (?,?,?,'unused','docx',?,?,'ready','now')",
                           (identifier, identifier + ".docx", identifier, text, len(text)))

    def tearDown(self):
        self.context.__exit__(None, None, None)
        self.environment.stop()
        self.directory.cleanup()

    def wait(self, identifier):
        for _ in range(300):
            response = self.client.get("/api/match/" + identifier)
            self.assertEqual(response.status_code, 200, response.text)
            task = response.json()
            if task["status"] in {"succeeded", "failed"}:
                return task
            time.sleep(.01)
        self.fail("matching did not terminate")

    def create(self, ids=None, jd=JD):
        response = self.client.post("/api/match", json={"jd_text": jd, "resume_ids": ids or ["a", "b", "c"]})
        self.assertEqual(response.status_code, 202, response.text)
        return response.json()["task_id"]

    def test_deterministic_sort_and_disabled_llm(self):
        with patch("app.matching.run_diagnose") as llm:
            first = self.wait(self.create(["b", "a", "c", "a"]))
            second = self.wait(self.create(["c", "a", "b"]))
        llm.assert_not_called()
        self.assertEqual(first, second)
        self.assertEqual(first["status"], "succeeded")
        self.assertEqual(first["progress"], 100)
        self.assertEqual([item["resume_id"] for item in first["items"]], ["a", "b", "c"])
        for item in first["items"]:
            self.assertTrue(item["degraded"])
            self.assertIsNone(item["reason"])
            self.assertNotIn("report", item)
            self.assertAlmostEqual(sum(part["weight"] for part in item["subscores"]), 1)
        self.assertIn("Python", first["items"][0]["matched"])
        self.assertIn("Python", first["items"][-1]["missing"])

    def test_validation_and_task_kind_isolation(self):
        cases = [({"jd_text": "短", "resume_ids": ["a"]}, 400),
                 ({"jd_text": "字" * 20001, "resume_ids": ["a"]}, 413),
                 ({"jd_text": JD, "resume_ids": []}, 400),
                 ({"jd_text": JD, "resume_ids": ["a"] * 51}, 400),
                 ({"jd_text": JD, "resume_ids": ["absent"]}, 404),
                 ({"jd_text": JD, "resume_ids": [4]}, 422), ({}, 422)]
        for payload, status in cases:
            response = self.client.post("/api/match", json=payload)
            self.assertEqual(response.status_code, status, response.text)
            self.assertTrue(response.json()["error"]["request_id"])
        app.state.store.update("resumes", "c", status="parsing")
        response = self.client.post("/api/match", json={"jd_text": JD, "resume_ids": ["c"]})
        self.assertEqual(response.status_code, 409)
        app.state.store.update("resumes", "c", status="ready", text="")
        self.assertEqual(self.client.post("/api/match", json={"jd_text": JD, "resume_ids": ["c"]}).status_code, 400)
        for path in ("/api/match/nope", "/api/match/nope/export"):
            self.assertEqual(self.client.get(path).status_code, 404)
        identifier = self.create(["a"])
        self.wait(identifier)
        self.assertEqual(self.client.get("/api/tasks/" + identifier).status_code, 404)
        with app.state.store.connect() as db:
            db.execute("INSERT INTO tasks (id,kind,status,created_at,updated_at) VALUES ('ingest','ingest','succeeded','now','now')")
        self.assertEqual(self.client.get("/api/match/ingest").status_code, 404)

    def test_ai_report_blending_and_persistent_cache(self):
        report = MOCK_RESPONSES["medium"].model_copy(deep=True)
        report.meta.mode = "llm"
        with patch.dict(os.environ, {"MATCH_LLM_TOP_K": "1"}), patch("app.matching.run_diagnose", new=AsyncMock(return_value=report)) as llm:
            first = self.wait(self.create())
            self.assertEqual(llm.await_count, 1)
            second = self.wait(self.create())
            self.assertEqual(llm.await_count, 1)
            self.assertEqual(first, second)
            refined = [item for item in first["items"] if not item["degraded"]]
            self.assertEqual(len(refined), 1)
            item = refined[0]
            DiagnoseResponse.model_validate(item["report"])
            self.assertEqual(item["score"], round(.35 * item["stage1_score"] + .65 * report.overall_score))
            self.assertEqual(item["reason"], report.summary)
            with app.state.store.connect() as db:
                self.assertEqual(db.execute("SELECT count(*) FROM llm_cache").fetchone()[0], 1)
                self.assertNotIn("jd_text", db.execute("SELECT payload FROM tasks LIMIT 1").fetchone()[0])

    def test_single_ai_failure_does_not_fail_batch(self):
        report = MOCK_RESPONSES["medium"].model_copy(deep=True)
        async def fake(resume, jd):
            if resume == TEXT:
                raise LlmError("timeout")
            return report
        with patch.dict(os.environ, {"MATCH_LLM_TOP_K": "8"}), patch("app.matching.run_diagnose", side_effect=fake):
            result = self.wait(self.create())
        self.assertEqual(result["status"], "succeeded")
        by_id = {item["resume_id"]: item for item in result["items"]}
        self.assertTrue(by_id["a"]["degraded"])
        self.assertTrue(by_id["b"]["degraded"])
        self.assertFalse(by_id["c"]["degraded"])

    def test_llm_concurrency_is_bounded_to_three(self):
        active = maximum = 0
        async def fake(resume, jd):
            nonlocal active, maximum
            active += 1
            maximum = max(maximum, active)
            try:
                await asyncio.sleep(.04)
                return MOCK_RESPONSES["medium"].model_copy(deep=True)
            finally:
                active -= 1
        with patch.dict(os.environ, {"MATCH_LLM_TOP_K": "8"}), patch("app.matching.run_diagnose", side_effect=fake):
            ids = [self.create(), self.create()]
            for identifier in ids:
                self.assertEqual(self.wait(identifier)["status"], "succeeded")
        self.assertLessEqual(maximum, 3)
        self.assertGreater(maximum, 1)

    def test_match_budget_cancels_ai_and_keeps_base_scores(self):
        cancelled = []
        async def stalled(resume, jd):
            try:
                await asyncio.sleep(10)
            finally:
                cancelled.append(True)
        with patch.dict(os.environ, {"MATCH_LLM_TOP_K": "8"}), \
             patch("app.matching.MATCH_TOTAL_TIMEOUT", .03), \
             patch("app.matching.run_diagnose", side_effect=stalled):
            result = self.wait(self.create())
        self.assertEqual(result["status"], "succeeded")
        self.assertEqual(len(cancelled), 3)
        self.assertTrue(all(item["degraded"] for item in result["items"]))
        self.assertTrue(all(item["score"] == item["stage1_score"] for item in result["items"]))

    def test_top_eight_limit_on_twenty_resumes(self):
        with app.state.store.connect() as db:
            for index in range(20):
                identifier = f"candidate-{index}"
                db.execute("INSERT INTO resumes (id,name,sha256,file_path,source,text,chars,status,created_at) "
                           "VALUES (?,?,?,'unused','docx',?,?,'ready','now')",
                           (identifier, identifier, identifier, TEXT + str(index), len(TEXT) + 2))
        report = MOCK_RESPONSES["medium"].model_copy(deep=True)
        with patch.dict(os.environ, {"MATCH_LLM_TOP_K": "99"}), \
             patch("app.matching.run_diagnose", new=AsyncMock(return_value=report)) as llm:
            result = self.wait(self.create([f"candidate-{index}" for index in range(20)]))
        self.assertEqual(llm.await_count, 8)
        self.assertEqual(len(result["items"]), 20)
        self.assertEqual(sum(not item["degraded"] for item in result["items"]), 8)

    def test_xlsx_export_handles_chinese_and_formula_names(self):
        app.state.store.update("resumes", "a", name="=HYPERLINK(危险)\x01.docx")
        identifier = self.create()
        self.wait(identifier)
        response = self.client.get(f"/api/match/{identifier}/export")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("filename*=UTF-8''", response.headers["content-disposition"])
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 4)
        self.assertTrue(sheet["B2"].value.startswith("=HYPERLINK"))
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertNotIn("\x01", sheet["B2"].value)
        with app.state.store.connect() as db:
            db.execute("INSERT INTO tasks (id,kind,status,created_at,updated_at) VALUES ('pending','match','running','now','now')")
        self.assertEqual(self.client.get("/api/match/pending/export").status_code, 409)

    def test_frontend_upload_to_matching_to_export(self):
        document = Document()
        document.add_paragraph(TEXT)
        stream = io.BytesIO()
        document.save(stream)
        upload = self.client.post("/api/resumes/batch", files=[("files", ("测试.docx", stream.getvalue()))])
        self.assertEqual(upload.status_code, 202)
        for _ in range(200):
            ingest = self.client.get("/api/tasks/" + upload.json()["task_id"]).json()
            if ingest["status"] == "succeeded":
                break
            time.sleep(.01)
        resume_id = ingest["items"][0]["resume_id"]
        self.assertEqual(self.client.get(f"/api/resumes/{resume_id}/text").json()["text"], TEXT)
        identifier = self.create([resume_id])
        result = self.wait(identifier)
        self.assertEqual(result["items"][0]["resume_id"], resume_id)
        self.assertEqual(self.client.get(f"/api/match/{identifier}/export").status_code, 200)

    def test_failed_match_and_restart_errors_are_strings(self):
        with patch("app.matching.score_resume", side_effect=RuntimeError("failure")):
            failed = self.wait(self.create())
        self.assertEqual(failed["status"], "failed")
        self.assertIsInstance(failed["error"], str)
        with app.state.store.connect() as db:
            db.execute("INSERT INTO tasks (id,kind,status,created_at,updated_at) VALUES ('restart','match','queued','now','now')")
        app.state.store.initialize()
        restarted = self.client.get("/api/match/restart").json()
        self.assertEqual(restarted["status"], "failed")
        self.assertIn("服务重启", restarted["error"])


class ScoringRulesTests(unittest.TestCase):
    def test_education_and_stable_terms(self):
        self.assertEqual(education_score("熟悉Python", "未写学历"), 100)
        self.assertEqual(education_score("本科以上", "未写学历"), 60)
        self.assertEqual(education_score("本科以上", "硕士"), 100)
        self.assertEqual(education_score("硕士以上", "本科"), 60)
        self.assertEqual(education_score("博士以上", "专科"), 30)
        self.assertEqual(jd_terms(JD), jd_terms(JD))
        self.assertNotEqual(cache_key(TEXT, JD), cache_key(TEXT, JD + " Linux"))

    def test_ocr_glued_english_terms(self):
        from app.matching import glued_runs, keyword_in_resume
        # OCR 把一行英文技能识别成连在一起的串
        text = "技能PythonFastAPIMySQLGitDocker\n项目使用FastAPI开发接口，部署到Linux"
        runs = glued_runs(text)
        for term in ("Python", "FastAPI", "MySQL", "Git", "Docker", "Linux"):
            self.assertTrue(keyword_in_resume(term, text, runs), term)
        self.assertFalse(keyword_in_resume("Redis", text, runs))
        # 普通单词不算粘连：JavaScript 里不能命中 Java，Go 太短也不按子串匹配
        normal = "熟悉 JavaScript 和 TypeScript，参与 Google 项目"
        self.assertEqual(glued_runs(normal), [])
        self.assertFalse(keyword_in_resume("Java", normal, glued_runs(normal)))
        self.assertFalse(keyword_in_resume("Go", text + "Google", glued_runs(text + "Google")))

    def test_diagnose_http_failure_is_normalized(self):
        with patch("app.routers.diagnose.diagnose", new=AsyncMock(side_effect=HTTPException(504))):
            with self.assertRaises(LlmError):
                asyncio.run(run_diagnose(TEXT, JD))


if __name__ == "__main__":
    unittest.main()
