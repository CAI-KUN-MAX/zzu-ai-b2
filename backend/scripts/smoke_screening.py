"""Local end-to-end smoke test with synthetic resumes and temporary storage.

Run from repository root:
    python backend/scripts/smoke_screening.py
    python backend/scripts/smoke_screening.py --live-llm

The second command sends fictional resume/JD text to the configured model.
"""
import argparse
import io
import json
import os
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
import pymupdf

from app.main import app

JD = "招聘Python后端开发实习生，本科及以上学历，熟悉Python、FastAPI和MySQL，掌握REST API设计，能够编写单元测试，使用Git协作，有Docker和Linux使用经验者优先。"
RESUME = "郑州大学计算机科学与技术专业本科，熟悉Python、FastAPI和MySQL，使用Git进行团队协作。负责校园报名系统的REST API开发，设计用户表和活动表，编写单元测试，使用Docker在Linux环境部署。"


def image_bytes(format):
    font = None
    for path in ("C:/Windows/Fonts/msyh.ttc", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if Path(path).exists():
            font = ImageFont.truetype(path, 32)
            break
    image = Image.new("RGB", (1200, 600), "white")
    draw = ImageDraw.Draw(image)
    for index, line in enumerate(("Python Backend Developer", "FastAPI MySQL Docker Linux Git",
                                  "Bachelor in Computer Science", "Develop REST APIs and unit tests")):
        draw.text((50, 50 + index * 110), line, font=font, fill="black")
    stream = io.BytesIO()
    image.save(stream, format=format)
    return stream.getvalue()


def fixtures():
    files = []
    for index in range(14):
        document = Document()
        document.add_paragraph(f"候选人{index + 1}。" + RESUME)
        stream = io.BytesIO()
        document.save(stream)
        files.append((f"candidate-{index + 1}.docx", stream.getvalue()))
    for index in range(3):
        with pymupdf.open() as document:
            page = document.new_page()
            page.insert_text((50, 80), f"Candidate {index + 15}: Python FastAPI MySQL backend developer. "
                             "Bachelor in Computer Science. Develop REST APIs and unit tests using Git Docker Linux.", fontsize=8)
            files.append((f"candidate-{index + 15}.pdf", document.tobytes()))
    png, jpg = image_bytes("PNG"), image_bytes("JPEG")
    with pymupdf.open() as document:
        page = document.new_page(width=600, height=300)
        page.insert_image(page.rect, stream=png)
        files.append(("candidate-18-scan.pdf", document.tobytes()))
    files.extend([("candidate-19.png", png), ("candidate-20.jpg", jpg)])
    return files


def wait(client, path, seconds=210):
    started = time.monotonic()
    while time.monotonic() - started < seconds:
        response = client.get(path)
        response.raise_for_status()
        task = response.json()
        if task["status"] in {"succeeded", "failed"}:
            assert task["status"] == "succeeded", task.get("error")
            return task
        time.sleep(.1)
    raise AssertionError("task timeout")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--live-llm", action="store_true")
    args = parser.parse_args()
    os.environ["OCR_PROVIDER"] = "rapidocr"
    os.environ["MATCH_LLM_TOP_K"] = "1" if args.live_llm else "0"
    os.environ["MOCK_MODE"] = "false"
    with tempfile.TemporaryDirectory() as directory:
        os.environ["DATA_DIR"] = directory
        with TestClient(app) as client:
            started = time.monotonic()
            response = client.post("/api/resumes/batch", files=[("files", fixture) for fixture in fixtures()])
            assert response.status_code == 202, response.text
            ingest = wait(client, "/api/tasks/" + response.json()["task_id"])
            assert len(ingest["items"]) == 20, len(ingest["items"])
            assert all(item["status"] in {"ready", "need_review"} for item in ingest["items"]), ingest
            parse_seconds = round(time.monotonic() - started, 2)
            request = {"jd_text": JD, "resume_ids": [item["resume_id"] for item in ingest["items"]]}
            started = time.monotonic()
            response = client.post("/api/match", json=request)
            assert response.status_code == 202, response.text
            identifier = response.json()["task_id"]
            matched = wait(client, "/api/match/" + identifier)
            match_seconds = round(time.monotonic() - started, 2)
            assert len(matched["items"]) == 20
            ai = [item for item in matched["items"] if not item["degraded"]]
            if args.live_llm:
                assert len(ai) == 1 and ai[0]["report"]["meta"]["mode"] == "llm", "Live model did not produce a validated report"
            exported = client.get(f"/api/match/{identifier}/export")
            assert exported.status_code == 200
            workbook = load_workbook(io.BytesIO(exported.content))
            assert workbook.active.max_row == 21
            started = time.monotonic()
            second = client.post("/api/match", json=request)
            cached = wait(client, "/api/match/" + second.json()["task_id"])
            assert cached["items"] == matched["items"]
            print(json.dumps({"resumes": 20, "parse_seconds": parse_seconds,
                              "match_seconds": match_seconds, "validated_ai_reports": len(ai),
                              "repeat_match_seconds": round(time.monotonic() - started, 2),
                              "export_rows": workbook.active.max_row - 1}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
